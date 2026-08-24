#!/usr/bin/env python3
"""Controlled artifact worker for generating project Office files.

The Electron process owns all project-root validation and path decisions.
This worker only accepts a JSON request on stdin and writes the requested
workbook to the already validated output path.
"""

from __future__ import annotations

import json
import math
import os
import re
import sys
import traceback
from typing import Any


MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 20_000
MAX_COLUMNS = 300
DEFAULT_FONT = "Microsoft YaHei"
HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
BORDER_COLOR = "D9E2EC"
TABLE_STYLE = "TableStyleMedium2"
ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F\uD800-\uDFFF]")


def emit(payload: dict[str, Any], status: int = 0) -> None:
    sys.stdout.write(json.dumps(payload, ensure_ascii=False))
    sys.stdout.write("\n")
    sys.stdout.flush()
    raise SystemExit(status)


def display_width(value: Any) -> int:
    if value is None:
        return 0
    text = str(value)
    width = 0
    for char in text:
        if char == "\n":
            width = max(width, 8)
            continue
        width += 1 if ord(char) < 128 else 2
    return width


def json_stringify(value: Any) -> str:
    try:
        return clean_text(json.dumps(value, ensure_ascii=False, separators=(",", ":")))
    except Exception:
        return clean_text(str(value))


def clean_text(value: str) -> str:
    return ILLEGAL_XML_RE.sub("", value)


def normalize_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return ""
        if isinstance(value, str) and value.startswith("="):
            return "'" + clean_text(value)
        if isinstance(value, str):
            return clean_text(value)
        return value
    if isinstance(value, (list, dict)):
        return json_stringify(value)
    return str(value)


def sanitize_sheet_name(name: Any, index: int, used: set[str]) -> str:
    raw = clean_text(str(name or "")).strip() or f"Sheet{index + 1}"
    cleaned = re.sub(r"[\\/?*\[\]:]", " ", raw).strip() or f"Sheet{index + 1}"
    cleaned = re.sub(r"\s+", " ", cleaned)
    base = cleaned[:31].strip() or f"Sheet{index + 1}"
    candidate = base
    suffix = 2
    while candidate.lower() in used:
        suffix_text = f" ({suffix})"
        candidate = f"{base[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used.add(candidate.lower())
    return candidate


def sanitize_table_name(sheet_name: str, index: int, used: set[str]) -> str:
    base = re.sub(r"[^A-Za-z0-9_]", "_", sheet_name).strip("_") or f"Table{index + 1}"
    if not re.match(r"^[A-Za-z_]", base):
        base = f"T_{base}"
    base = base[:200]
    candidate = base
    suffix = 2
    while candidate.lower() in used:
        candidate = f"{base[:190]}_{suffix}"
        suffix += 1
    used.add(candidate.lower())
    return candidate


def build_matrix(rows: list[Any]) -> tuple[list[list[Any]], bool]:
    if not rows:
        return [["（无数据）"]], False

    all_arrays = all(isinstance(row, list) for row in rows)
    if all_arrays:
        max_cols = min(max((len(row) for row in rows), default=1), MAX_COLUMNS)
        matrix: list[list[Any]] = []
        for row in rows:
            cells = list(row[:max_cols])
            cells.extend([""] * (max_cols - len(cells)))
            matrix.append([normalize_cell(cell) for cell in cells])
        return matrix, True

    keys: list[str] = []
    seen: set[str] = set()
    normalized_rows: list[dict[str, Any]] = []
    max_array_cols = 0
    for raw in rows:
        if isinstance(raw, dict):
            record = raw
        elif isinstance(raw, list):
            max_array_cols = max(max_array_cols, len(raw))
            record = {f"列{index + 1}": value for index, value in enumerate(raw)}
        else:
            record = {"value": raw}
        normalized_rows.append(record)
        for key in record.keys():
            key_text = clean_text(str(key or "")).strip() or "value"
            if key_text not in seen:
                seen.add(key_text)
                keys.append(key_text)
            if len(keys) >= MAX_COLUMNS:
                break

    if max_array_cols:
        for index in range(max_array_cols):
            key = f"列{index + 1}"
            if key not in seen and len(keys) < MAX_COLUMNS:
                seen.add(key)
                keys.append(key)

    if not keys:
        keys = ["value"]

    matrix = [keys]
    for record in normalized_rows:
        matrix.append([normalize_cell(record.get(key, "")) for key in keys])
    return matrix, True


def header_is_table_safe(matrix: list[list[Any]]) -> bool:
    if len(matrix) < 2 or not matrix[0]:
        return False
    seen: set[str] = set()
    for cell in matrix[0]:
        text = str(cell or "").strip()
        if not text:
            return False
        key = text.lower()
        if key in seen:
            return False
        seen.add(key)
    return True


def infer_number_format(header: str, value: Any) -> str | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    text = header.lower()
    chinese_header = header
    if any(token in text for token in ["amount", "cost", "price", "total", "money"]) or any(
        token in chinese_header for token in ["金额", "造价", "费用", "单价", "合价", "总价"]
    ):
        return "#,##0.00"
    if any(token in text for token in ["percent", "rate", "ratio"]) or any(
        token in chinese_header for token in ["百分比", "比例", "费率", "税率"]
    ):
        return "0.00%" if abs(float(value)) <= 1 else "0.00"
    if any(token in text for token in ["qty", "quantity", "area", "volume", "length", "weight"]) or any(
        token in chinese_header for token in ["数量", "工程量", "面积", "体积", "长度", "重量", "厚度", "高度", "宽度"]
    ):
        return "#,##0.000"
    if isinstance(value, int):
        return "#,##0"
    return "#,##0.00"


def apply_basic_style(ws: Any, matrix: list[list[Any]], table_safe: bool, table_name: str | None) -> dict[str, int]:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    max_row = ws.max_row
    max_col = ws.max_column
    thin_side = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(name=DEFAULT_FONT, bold=True, color=HEADER_FONT, size=10)
    body_font = Font(name=DEFAULT_FONT, size=10)

    ws.sheet_view.showGridLines = False

    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col), start=1):
        row_text_width = 0
        for cell in row:
            cell.border = border
            cell.font = header_font if row_index == 1 else body_font
            if row_index == 1:
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                value = cell.value
                header = str(ws.cell(row=1, column=cell.column).value or "")
                fmt = infer_number_format(header, value)
                if fmt:
                    cell.number_format = fmt
                    horizontal = "right"
                else:
                    horizontal = "left"
                text_width = display_width(value)
                row_text_width = max(row_text_width, text_width)
                cell.alignment = Alignment(
                    horizontal=horizontal,
                    vertical="center",
                    wrap_text=text_width > 26 or ("\n" in str(value)),
                )
        if row_index == 1:
            ws.row_dimensions[row_index].height = 24
        elif row_text_width > 80:
            ws.row_dimensions[row_index].height = 54
        elif row_text_width > 45:
            ws.row_dimensions[row_index].height = 36
        else:
            ws.row_dimensions[row_index].height = 20

    for col_index in range(1, max_col + 1):
        col_letter = get_column_letter(col_index)
        values = [ws.cell(row=row_index, column=col_index).value for row_index in range(1, min(max_row, 200) + 1)]
        widths = [display_width(value) for value in values]
        header_text = str(ws.cell(row=1, column=col_index).value or "")
        numeric_count = sum(
            1
            for value in values[1:]
            if isinstance(value, (int, float)) and not isinstance(value, bool)
        )
        max_width = max(widths or [8])
        if numeric_count >= max(1, len(values[1:]) // 2):
            width = min(max(max_width + 2, 10), 18)
        elif max_width > 36 or any("\n" in str(value) for value in values):
            width = min(max(max_width * 0.85 + 4, display_width(header_text) + 2, 16), 42)
        else:
            width = min(max(max_width + 2, 10), 30)
        ws.column_dimensions[col_letter].width = width

    if max_row > 1:
        ws.freeze_panes = "A2"
    if max_row >= 1 and max_col >= 1:
        ws.auto_filter.ref = ws.dimensions

    if table_safe and table_name and max_row >= 2 and max_col >= 1:
        table = Table(displayName=table_name, ref=ws.dimensions)
        style = TableStyleInfo(
            name=TABLE_STYLE,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = style
        ws.add_table(table)

    return {"rows": max_row, "columns": max_col}


def write_excel(payload: dict[str, Any]) -> dict[str, Any]:
    from openpyxl import Workbook

    output_path = str(payload.get("output_path") or "").strip()
    if not output_path:
        raise ValueError("missing output_path")

    sheets = payload.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("Excel 产物至少需要一个 sheet。")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    warnings: list[str] = []
    used_sheet_names: set[str] = set()
    used_table_names: set[str] = set()
    sheet_summaries: list[dict[str, Any]] = []

    for index, sheet in enumerate(sheets[:MAX_SHEETS]):
        if not isinstance(sheet, dict):
            sheet = {"name": f"Sheet{index + 1}", "rows": []}
        rows = sheet.get("rows") if isinstance(sheet.get("rows"), list) else []
        if len(rows) > MAX_ROWS_PER_SHEET:
            raise ValueError(f"Sheet {sheet.get('name') or index + 1} 行数超过 {MAX_ROWS_PER_SHEET}。")

        sheet_name = sanitize_sheet_name(sheet.get("name"), index, used_sheet_names)
        matrix, has_header = build_matrix(rows)
        ws = workbook.create_sheet(sheet_name)
        for row in matrix:
            ws.append(row)

        table_safe = has_header and header_is_table_safe(matrix)
        table_name = sanitize_table_name(sheet_name, index, used_table_names) if table_safe else None
        summary = apply_basic_style(ws, matrix, table_safe, table_name)
        summary["name"] = sheet_name
        sheet_summaries.append(summary)
        if not rows:
            warnings.append(f"{sheet_name} 为空表，已生成占位行。")

    if len(sheets) > MAX_SHEETS:
        warnings.append(f"仅写入前 {MAX_SHEETS} 个 sheet。")

    metadata = payload.get("metadata")
    if isinstance(metadata, dict) and metadata:
        sheet_name = sanitize_sheet_name("metadata", len(sheet_summaries), used_sheet_names)
        ws = workbook.create_sheet(sheet_name)
        ws.append(["字段", "值"])
        for key, value in metadata.items():
            ws.append([clean_text(str(key)), normalize_cell(value)])
        summary = apply_basic_style(
            ws,
            [["字段", "值"], *[[clean_text(str(key)), value] for key, value in metadata.items()]],
            True,
            sanitize_table_name(sheet_name, len(sheet_summaries), used_table_names),
        )
        summary["name"] = sheet_name
        sheet_summaries.append(summary)

    workbook.save(output_path)
    size_bytes = os.path.getsize(output_path)
    return {
        "success": True,
        "output_path": output_path,
        "size_bytes": size_bytes,
        "sheet_count": len(sheet_summaries),
        "sheets": sheet_summaries,
        "warnings": warnings,
    }


def text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json_stringify(value)
    return clean_text(str(value))


def apply_docx_font(run: Any, size: int = 10, bold: bool = False, color: str | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = DEFAULT_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_docx_table(table: Any) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), HEADER_FILL)
                cell._tc.get_or_add_tcPr().append(shading)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    apply_docx_font(run, 9, row_index == 0, HEADER_FONT if row_index == 0 else None)


def write_docx(payload: dict[str, Any]) -> dict[str, Any]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    output_path = str(payload.get("output_path") or "").strip()
    title = text_value(payload.get("title")).strip()
    blocks = payload.get("blocks")
    if not output_path:
        raise ValueError("missing output_path")
    if not title:
        raise ValueError("DOCX 产物必须提供 title。")
    if not isinstance(blocks, list) or not blocks:
        raise ValueError("DOCX 产物至少需要一个 block。")
    if len(blocks) > 300:
        raise ValueError("DOCX blocks 不能超过 300 个。")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    document = Document()
    for section in document.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = DEFAULT_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)
    normal.font.size = Pt(10.5)
    for level in range(1, 4):
        style = styles[f"Heading {level}"]
        style.font.name = DEFAULT_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)
        style.font.color.rgb = RGBColor.from_string(HEADER_FILL)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_docx_font(title_paragraph.add_run(title), 20, True, HEADER_FILL)
    subtitle = text_value(payload.get("subtitle")).strip()
    if subtitle:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        apply_docx_font(paragraph.add_run(subtitle), 11, False, "64748B")

    metadata = payload.get("metadata")
    if isinstance(metadata, dict) and metadata:
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "字段"
        table.rows[0].cells[1].text = "内容"
        for key, value in metadata.items():
            cells = table.add_row().cells
            cells[0].text = text_value(key)
            cells[1].text = text_value(value)
        style_docx_table(table)
        document.add_paragraph()

    warnings: list[str] = []
    for index, raw_block in enumerate(blocks):
        block = raw_block if isinstance(raw_block, dict) else {"type": "paragraph", "text": raw_block}
        block_type = str(block.get("type") or "paragraph").strip().lower()
        if block_type == "heading":
            level = int(block.get("level") or 1)
            level = max(1, min(level, 3))
            text = text_value(block.get("text")).strip()
            if text:
                document.add_heading(text, level=level)
        elif block_type == "paragraph":
            text = text_value(block.get("text")).strip()
            if text:
                document.add_paragraph(text)
        elif block_type == "bullets":
            items = block.get("items") if isinstance(block.get("items"), list) else []
            for item in items[:100]:
                document.add_paragraph(text_value(item), style="List Bullet")
        elif block_type == "key_values":
            items = block.get("items") if isinstance(block.get("items"), list) else []
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "字段"
            table.rows[0].cells[1].text = "内容"
            for item in items[:200]:
                record = item if isinstance(item, dict) else {"key": "", "value": item}
                cells = table.add_row().cells
                cells[0].text = text_value(record.get("key"))
                cells[1].text = text_value(record.get("value"))
            style_docx_table(table)
        elif block_type == "table":
            columns = [text_value(item).strip() for item in (block.get("columns") or [])][:20]
            rows = block.get("rows") if isinstance(block.get("rows"), list) else []
            if not columns and rows and isinstance(rows[0], dict):
                columns = [text_value(item).strip() for item in rows[0].keys()][:20]
            if not columns:
                warnings.append(f"block {index + 1} 表格缺少 columns，已跳过。")
                continue
            table = document.add_table(rows=1, cols=len(columns))
            for column_index, column in enumerate(columns):
                table.rows[0].cells[column_index].text = column
            for raw_row in rows[:2000]:
                row = raw_row if isinstance(raw_row, dict) else {}
                cells = table.add_row().cells
                for column_index, column in enumerate(columns):
                    cells[column_index].text = text_value(row.get(column))
            style_docx_table(table)
        elif block_type == "note":
            tone = str(block.get("tone") or "note").lower()
            prefix = "⚠ 待确认：" if tone == "warning" else "说明："
            paragraph = document.add_paragraph()
            apply_docx_font(paragraph.add_run(prefix), 10, True, "B45309" if tone == "warning" else HEADER_FILL)
            apply_docx_font(paragraph.add_run(text_value(block.get("text"))), 10)
        else:
            warnings.append(f"不支持的 DOCX block 类型 {block_type}，已跳过。")

    document.save(output_path)
    return {
        "success": True,
        "output_path": output_path,
        "size_bytes": os.path.getsize(output_path),
        "block_count": len(blocks),
        "warnings": warnings,
    }


def set_pptx_text_frame(text_frame: Any, lines: list[str], size: int, color: tuple[int, int, int], bold: bool = False) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    text_frame.clear()
    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = clean_text(line)
        paragraph.font.name = DEFAULT_FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor(*color)


def add_pptx_title(slide: Any, title: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(11.9), Inches(0.7))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.name = DEFAULT_FONT
    paragraph.font.size = Pt(26)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(31, 78, 120)


def add_pptx_footer(slide: Any, text: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(0.7), Inches(7.12), Inches(11.9), Inches(0.22))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = DEFAULT_FONT
    paragraph.font.size = Pt(8)
    paragraph.font.color.rgb = RGBColor(100, 116, 139)


def write_pptx(payload: dict[str, Any]) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    output_path = str(payload.get("output_path") or "").strip()
    deck_title = text_value(payload.get("title")).strip()
    slides = payload.get("slides")
    if not output_path:
        raise ValueError("missing output_path")
    if not deck_title:
        raise ValueError("PPTX 产物必须提供 title。")
    if not isinstance(slides, list) or not slides:
        raise ValueError("PPTX 产物至少需要一页 slide。")
    if len(slides) > 80:
        raise ValueError("PPTX slides 不能超过 80 页。")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]
    warnings: list[str] = []

    normalized_slides = list(slides)
    first_type = str(normalized_slides[0].get("type") if isinstance(normalized_slides[0], dict) else "")
    if first_type != "title":
        normalized_slides.insert(0, {
            "type": "title",
            "title": deck_title,
            "subtitle": text_value(payload.get("subtitle")),
        })

    for slide_index, raw_slide in enumerate(normalized_slides):
        data = raw_slide if isinstance(raw_slide, dict) else {"type": "content", "title": f"第 {slide_index + 1} 页", "body": raw_slide}
        slide_type = str(data.get("type") or "content").strip().lower()
        title = text_value(data.get("title") or deck_title).strip()
        slide = presentation.slides.add_slide(blank_layout)

        if slide_type in {"title", "section", "closing"}:
            accent = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
            accent.fill.solid()
            accent.fill.fore_color.rgb = RGBColor(31, 78, 120) if slide_type != "closing" else RGBColor(15, 23, 42)
            accent.line.fill.background()
            title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.15), Inches(11.3), Inches(1.3))
            set_pptx_text_frame(title_box.text_frame, [title], 32 if slide_type == "title" else 30, (255, 255, 255), True)
            subtitle = text_value(data.get("subtitle") or data.get("body")).strip()
            if subtitle:
                subtitle_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.65), Inches(11.3), Inches(1.1))
                set_pptx_text_frame(subtitle_box.text_frame, [subtitle], 16, (219, 234, 254))
        else:
            add_pptx_title(slide, title)
            if slide_type == "bullets":
                bullets = [text_value(item) for item in (data.get("bullets") or [])][:12]
                box = slide.shapes.add_textbox(Inches(1.0), Inches(1.45), Inches(11.2), Inches(5.2))
                frame = box.text_frame
                frame.clear()
                for bullet_index, bullet in enumerate(bullets):
                    paragraph = frame.paragraphs[0] if bullet_index == 0 else frame.add_paragraph()
                    paragraph.text = bullet
                    paragraph.level = 0
                    paragraph.font.name = DEFAULT_FONT
                    paragraph.font.size = Pt(20)
                    paragraph.font.color.rgb = RGBColor(30, 41, 59)
                    paragraph.space_after = Pt(10)
            elif slide_type == "table":
                columns = [text_value(item).strip() for item in (data.get("columns") or [])][:12]
                rows = data.get("rows") if isinstance(data.get("rows"), list) else []
                if not columns and rows and isinstance(rows[0], dict):
                    columns = [text_value(item).strip() for item in rows[0].keys()][:12]
                if columns:
                    visible_rows = rows[:20]
                    shape = slide.shapes.add_table(len(visible_rows) + 1, len(columns), Inches(0.8), Inches(1.45), Inches(11.75), Inches(5.1))
                    table = shape.table
                    for column_index, column in enumerate(columns):
                        table.cell(0, column_index).text = column
                    for row_index, raw_row in enumerate(visible_rows, start=1):
                        row = raw_row if isinstance(raw_row, dict) else {}
                        for column_index, column in enumerate(columns):
                            table.cell(row_index, column_index).text = text_value(row.get(column))
                    for row_index in range(len(visible_rows) + 1):
                        for column_index in range(len(columns)):
                            cell = table.cell(row_index, column_index)
                            cell.fill.solid()
                            cell.fill.fore_color.rgb = RGBColor(31, 78, 120) if row_index == 0 else RGBColor(248, 250, 252)
                            for paragraph in cell.text_frame.paragraphs:
                                paragraph.font.name = DEFAULT_FONT
                                paragraph.font.size = Pt(10)
                                paragraph.font.bold = row_index == 0
                                paragraph.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else RGBColor(30, 41, 59)
                    if len(rows) > 20:
                        warnings.append(f"{title} 表格仅展示前 20 行。")
                else:
                    warnings.append(f"{title} 表格缺少 columns，已生成空内容页。")
            else:
                body = text_value(data.get("body")).strip()
                if body:
                    box = slide.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(11.2), Inches(4.9))
                    set_pptx_text_frame(box.text_frame, [body], 19, (30, 41, 59))
            add_pptx_footer(slide, f"{deck_title}  ·  {slide_index + 1}")

        notes = text_value(data.get("notes")).strip()
        if notes:
            try:
                slide.notes_slide.notes_text_frame.text = notes
            except Exception:
                warnings.append(f"{title} 的 notes 未能写入。")

    presentation.save(output_path)
    return {
        "success": True,
        "output_path": output_path,
        "size_bytes": os.path.getsize(output_path),
        "slide_count": len(normalized_slides),
        "warnings": warnings,
    }


def main() -> None:
    try:
        raw = sys.stdin.read()
        payload = json.loads(raw or "{}")
        action = payload.get("action")
        if action == "write_excel":
            result = write_excel(payload)
        elif action == "write_docx":
            result = write_docx(payload)
        elif action == "write_pptx":
            result = write_pptx(payload)
        else:
            raise ValueError(f"unsupported action: {action}")
        emit(result, 0)
    except SystemExit:
        raise
    except Exception as exc:
        emit(
            {
                "success": False,
                "error": str(exc),
                "traceback": traceback.format_exc(limit=8),
            },
            1,
        )


if __name__ == "__main__":
    main()
